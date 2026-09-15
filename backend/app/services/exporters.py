"""
Renders a deck (the same slide schema used by the in-browser editor) into
a real, downloadable PPTX or PDF file.

Chart slides: if the frontend captured a real chart image (see
chartCapture.tsx, which screenshots the live-rendered chart via
html2canvas), that PNG is embedded directly, so the export shows the
actual bar/line/pie graphic, not just its numbers. If no captured image is
available for a given chart slide (e.g. an older client, or the capture
failed), it falls back to rendering the same data as a plain table so the
export never silently drops the content.
"""
import base64
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor, black
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

DEFAULT_ACCENT = "1a1a1a"


def _data_url_to_bytes(data_url: str) -> bytes | None:
    if not data_url or "," not in data_url:
        return None
    try:
        return base64.b64decode(data_url.split(",", 1)[1])
    except Exception:
        return None


def _hex_or_default(color: str | None, default: str) -> str:
    if color and color.startswith("#") and len(color) in (4, 7):
        return color.lstrip("#")
    return default


def build_pptx(deck_name: str, slides: list[dict], theme: dict, chart_images: dict[str, str] | None = None) -> bytes:
    chart_images = chart_images or {}
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.63)
    blank_layout = prs.slide_layouts[6]

    accent_hex = _hex_or_default(theme.get("primary_color"), DEFAULT_ACCENT)
    accent_rgb = RGBColor.from_string(accent_hex.upper())

    for slide_data in slides:
        slide = prs.slides.add_slide(blank_layout)
        _render_pptx_slide(slide, slide_data, accent_rgb, prs.slide_width, prs.slide_height, chart_images)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_title(slide, text: str, accent: RGBColor, width, top=Inches(0.4)):
    box = slide.shapes.add_textbox(Inches(0.5), top, width - Inches(1), Inches(0.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = accent
    return box


def _render_pptx_slide(slide, data: dict, accent: RGBColor, slide_width, slide_height, chart_images: dict[str, str]):
    slide_type = data.get("type")
    title = data.get("title", "")

    if slide_type == "cover":
        img_bytes = _data_url_to_bytes(data.get("image_data_url") or "")
        top = Inches(1.0)
        if img_bytes:
            slide.shapes.add_picture(io.BytesIO(img_bytes), Inches(4.2), Inches(0.5), height=Inches(1.2))
            top = Inches(1.9)
        box = slide.shapes.add_textbox(Inches(0.5), top, slide_width - Inches(1), Inches(1))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = accent
        p.alignment = 2
        for block in data.get("blocks", []):
            sub = tf.add_paragraph()
            sub.text = block.get("text", "")
            sub.font.size = Pt(16)
            sub.alignment = 2
        return

    if slide_type == "team":
        _add_title(slide, title, accent, slide_width)
        members = data.get("team_members") or []
        x = Inches(0.5)
        col_width = (slide_width - Inches(1)) / max(len(members), 1)
        for m in members:
            box = slide.shapes.add_textbox(x, Inches(1.4), col_width - Inches(0.2), Inches(3.5))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = m.get("name", "")
            p.font.bold = True
            p.font.size = Pt(14)
            role_p = tf.add_paragraph()
            role_p.text = m.get("role", "")
            role_p.font.size = Pt(11)
            role_p.font.italic = True
            bio_p = tf.add_paragraph()
            bio_p.text = m.get("bio", "")
            bio_p.font.size = Pt(10)
            x += col_width
        return

    if slide_type == "toc":
        _add_title(slide, title, accent, slide_width)
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), slide_width - Inches(1.4), Inches(3.8))
        tf = box.text_frame
        tf.word_wrap = True
        entries = data.get("toc_entries") or []
        for i, entry in enumerate(entries):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"{i + 2}.  {entry.get('title', '')}"
            p.font.size = Pt(16)
        return

    if slide_type == "image":
        _add_title(slide, title, accent, slide_width)
        img_bytes = _data_url_to_bytes(data.get("image_data_url") or "")
        if img_bytes:
            slide.shapes.add_picture(io.BytesIO(img_bytes), Inches(1.5), Inches(1.3), width=Inches(7))
        return

    if slide_type == "chart":
        _add_title(slide, title, accent, slide_width)
        captured = chart_images.get(data.get("id", ""))
        captured_bytes = _data_url_to_bytes(captured) if captured else None
        if captured_bytes:
            slide.shapes.add_picture(io.BytesIO(captured_bytes), Inches(1.0), Inches(1.2), width=Inches(8))
            return
        # Fallback: no captured image for this chart, render the underlying
        # data as a table so the export still shows something accurate.
        chart = data.get("chart") or {}
        columns = chart.get("columns", [])
        rows = chart.get("rows", [])
        note_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), slide_width - Inches(1), Inches(0.3))
        note_box.text_frame.text = "Shown as a data table (chart image unavailable for this export)."
        note_box.text_frame.paragraphs[0].font.size = Pt(9)
        note_box.text_frame.paragraphs[0].font.italic = True
        if columns and rows:
            table_shape = slide.shapes.add_table(
                len(rows) + 1, len(columns), Inches(0.5), Inches(1.5), slide_width - Inches(1), Inches(3)
            )
            table = table_shape.table
            for ci, col in enumerate(columns):
                table.cell(0, ci).text = col
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    table.cell(ri + 1, ci).text = str(cell)
        return

    _add_title(slide, title, accent, slide_width)
    img_bytes = _data_url_to_bytes(data.get("image_data_url") or "")
    body_top = Inches(1.3)
    if img_bytes:
        slide.shapes.add_picture(io.BytesIO(img_bytes), Inches(2.5), Inches(1.2), height=Inches(1.4))
        body_top = Inches(2.8)
    box = slide.shapes.add_textbox(Inches(0.6), body_top, slide_width - Inches(1.2), slide_height - body_top - Inches(0.3))
    tf = box.text_frame
    tf.word_wrap = True
    blocks = data.get("blocks", [])
    for i, block in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = block.get("text", "")
        p.font.size = Pt(15)
        p.space_after = Pt(10)


def build_pdf(deck_name: str, slides: list[dict], theme: dict, chart_images: dict[str, str] | None = None) -> bytes:
    chart_images = chart_images or {}
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    styles = getSampleStyleSheet()

    accent_hex = _hex_or_default(theme.get("primary_color"), DEFAULT_ACCENT)
    accent_color = HexColor(f"#{accent_hex}")

    title_style = ParagraphStyle("SlideTitle", parent=styles["Heading1"], textColor=accent_color, spaceAfter=14)
    body_style = ParagraphStyle("SlideBody", parent=styles["BodyText"], fontSize=12, leading=17, spaceAfter=10)
    source_style = ParagraphStyle("Source", parent=styles["BodyText"], fontSize=8, textColor=black, spaceAfter=14)
    subtitle_style = ParagraphStyle("Subtitle", parent=styles["BodyText"], fontSize=14, alignment=1, spaceAfter=8)

    story = []
    for i, data in enumerate(slides):
        if i > 0:
            story.append(PageBreak())
        story.extend(_render_pdf_slide(data, title_style, body_style, source_style, subtitle_style, chart_images))

    doc.build(story)
    return buf.getvalue()


def _render_pdf_slide(data: dict, title_style, body_style, source_style, subtitle_style, chart_images: dict[str, str]) -> list:
    slide_type = data.get("type")
    title = data.get("title", "")
    story: list = []

    img_bytes = _data_url_to_bytes(data.get("image_data_url") or "")

    if slide_type == "cover":
        if img_bytes:
            try:
                story.append(RLImage(io.BytesIO(img_bytes), width=1.2 * inch, height=1.2 * inch))
                story.append(Spacer(1, 12))
            except Exception:
                pass
        story.append(Paragraph(title, ParagraphStyle("CoverTitle", parent=title_style, alignment=1, fontSize=26)))
        for block in data.get("blocks", []):
            story.append(Paragraph(block.get("text", ""), subtitle_style))
        return story

    story.append(Paragraph(title, title_style))

    if slide_type == "team":
        for m in data.get("team_members") or []:
            story.append(Paragraph(f"<b>{m.get('name', '')}</b> &mdash; {m.get('role', '')}", body_style))
            story.append(Paragraph(m.get("bio", ""), source_style))
        return story

    if slide_type == "toc":
        for i, entry in enumerate(data.get("toc_entries") or []):
            story.append(Paragraph(f"{i + 2}. {entry.get('title', '')}", body_style))
        return story

    if slide_type == "image":
        if img_bytes:
            try:
                story.append(RLImage(io.BytesIO(img_bytes), width=5 * inch, height=3 * inch))
            except Exception:
                story.append(Paragraph("(image could not be rendered)", source_style))
        return story

    if slide_type == "chart":
        captured = chart_images.get(data.get("id", ""))
        captured_bytes = _data_url_to_bytes(captured) if captured else None
        if captured_bytes:
            try:
                story.append(RLImage(io.BytesIO(captured_bytes), width=5.5 * inch, height=3.1 * inch))
                return story
            except Exception:
                pass
        chart = data.get("chart") or {}
        columns = chart.get("columns", [])
        rows = chart.get("rows", [])
        story.append(Paragraph("Shown as a data table (chart image unavailable for this export).", source_style))
        if columns and rows:
            table_data = [columns] + rows
            table = Table(table_data, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), HexColor("#eeeeee")),
                ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#cccccc")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]))
            story.append(table)
        return story

    if img_bytes:
        try:
            story.append(RLImage(io.BytesIO(img_bytes), width=4 * inch, height=2.2 * inch))
            story.append(Spacer(1, 10))
        except Exception:
            pass
    for block in data.get("blocks", []):
        story.append(Paragraph(block.get("text", ""), body_style))
        source = block.get("source", "")
        if source:
            story.append(Paragraph(f"Source: {source}", source_style))
    return story
